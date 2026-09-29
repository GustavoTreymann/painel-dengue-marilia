"""Monta o Relatorio Parcial sobre o modelo oficial da Univesp.
Uso (na raiz do repositorio): python docs/relatorio/gerar_relatorio.py
Depois rode atualizar_campos.ps1 para o Word numerar sumario, listas e folhas."""
import copy
import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RAIZ = Path(__file__).resolve().parents[2]
FIG = RAIZ / "docs" / "relatorio" / "figuras"
MODELO = RAIZ / "docs" / "modelos" / "Modelo_Relatorio_Parcial.docx"
SAIDA = RAIZ / "docs" / "relatorio" / "Relatorio_Parcial_PI2.docx"

AUTOR = "Luiz Gustavo Jacobuce Treymann"
TITULO = "Painel web de vigilância da dengue em Marília: dados abertos, clima e acessibilidade"
CURSO = "Bacharelado em Ciência de Dados"
TUTORA = "Jessica Caroline Pena Alves Da Silva"
FONTE_AUTOR = "Fonte: elaborado pelo autor (2026)."

d = docx.Document(MODELO)
corpo = d.element.body
filhos = list(corpo.iterchildren())

BASE, BULLET, CITACAO = "a) texto-base", "b) texto com bullets", "c) citação"
H1, H2, RESREF = "1) título nivel 1", "2) título nivel 2", "f) Resumo/Referências"


# ---------------------------------------------------------------- utilitarios
def trocar_texto(p_el, texto):
    """Mantem a formatacao do primeiro run e substitui o texto do paragrafo."""
    runs = p_el.findall(qn("w:r"))
    for r in runs[1:]:
        p_el.remove(r)
    if not runs:
        r = OxmlElement("w:r"); p_el.append(r); runs = [r]
    for t in runs[0].findall(qn("w:t")):
        runs[0].remove(t)
    t = OxmlElement("w:t"); t.text = texto; t.set(qn("xml:space"), "preserve")
    runs[0].append(t)


def campo(paragrafo, instrucao, marcador="1"):
    """Insere um campo do Word (SEQ, PAGE, TOC) marcado para atualizacao."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instrucao)
    fld.set(qn("w:dirty"), "true")
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = marcador
    r.append(t); fld.append(r)
    paragrafo._p.append(fld)


ANCORA = filhos[-1]  # sectPr final: todo conteudo novo entra antes dele


def novo_par(estilo=BASE):
    p = d.add_paragraph(style=estilo)
    ANCORA.addprevious(p._p)
    return p


def texto_rico(p, texto):
    """Aceita **negrito** e *italico* simples dentro do texto."""
    for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not parte:
            continue
        if parte.startswith("**"):
            p.add_run(parte[2:-2]).bold = True
        elif parte.startswith("*"):
            p.add_run(parte[1:-1]).italic = True
        else:
            p.add_run(parte)
    return p


def P(texto, estilo=BASE):
    return texto_rico(novo_par(estilo), texto)


def B(*itens):
    for item in itens:
        P(item, BULLET)


def legenda(tipo, titulo):
    p = novo_par(BASE)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if tipo == "Figura" else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f"{tipo} "); r.font.size = Pt(10)
    campo(p, f" SEQ {tipo} \\* ARABIC ")
    p.runs  # noqa
    r = p.add_run(f" – {titulo}"); r.font.size = Pt(10)
    for fld in p._p.findall(qn("w:fldSimple")):
        for rr in fld.findall(qn("w:r")):
            rpr = OxmlElement("w:rPr"); sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "20")
            rpr.append(sz); rr.insert(0, rpr)
    return p


def fonte(texto=FONTE_AUTOR):
    p = novo_par(BASE)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto); r.font.size = Pt(10)


def figura(arquivo, titulo, largura_cm=15.5, origem=FONTE_AUTOR):
    legenda("Figura", titulo)
    p = novo_par(BASE)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / arquivo), width=Cm(largura_cm))
    fonte(origem)


def borda(celula, lados):
    tcPr = celula._tc.get_or_add_tcPr()
    bordas = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{lado}")
        if lado in lados:
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "8"); b.set(qn("w:color"), "000000")
        else:
            b.set(qn("w:val"), "nil")
        bordas.append(b)
    tcPr.append(bordas)


def tabela(titulo, cabecalho, linhas, larguras_cm, origem=FONTE_AUTOR):
    """Tabela no padrao IBGE/ABNT: bordas horizontais, laterais abertas."""
    legenda("Tabela", titulo)
    t = d.add_table(rows=1 + len(linhas), cols=len(cabecalho))
    ANCORA.addprevious(t._tbl)
    t.alignment = 1
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
    for i, linha in enumerate([cabecalho] + linhas):
        for j, valor in enumerate(linha):
            c = t.cell(i, j)
            c.width = Cm(larguras_cm[j])
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r = p.add_run(str(valor)); r.font.size = Pt(10); r.font.name = "Times New Roman"
            if i == 0:
                r.bold = True
            if i < len(linhas):
                p.paragraph_format.keep_with_next = True
            lados = []
            if i == 0:
                lados = ["top", "bottom"]
            if i == len(linhas):
                lados.append("bottom")
            borda(c, lados)
    fonte(origem)


# ------------------------------------------------------------------- capa
capa = {1: "", 8: AUTOR, 9: "", 18: TITULO, 19: "", 41: "Marília - SP", 42: "2026",
        44: "", 56: TITULO, 57: "", 80: "Marília - SP", 81: "2026"}
for i, texto in capa.items():
    trocar_texto(filhos[i], texto)
# O titulo ocupa duas linhas: duas linhas em branco a menos por capa evitam que o ano pule de pagina.
for i in (38, 39, 77, 78):
    corpo.remove(filhos[i])
# Folha de rosto sempre em pagina propria, sem depender do preenchimento da capa.
filhos[43].get_or_add_pPr().append(OxmlElement('w:pageBreakBefore'))
trocar_texto(filhos[64],
             "Relatório Técnico-Científico apresentado na disciplina de Projeto Integrador em Computação II "
             f"para o curso de {CURSO} da Universidade Virtual do Estado de São Paulo (UNIVESP).")

trocar_texto(filhos[83],
             f"TREYMANN, Luiz Gustavo Jacobuce. {TITULO}. {{{{FOLHAS}}}}f. Relatório Técnico-Científico. "
             f"{CURSO} – Universidade Virtual do Estado de São Paulo. Tutora: {TUTORA}. Polo Marília, 2026.")

RESUMO = (
    "A dengue é um problema recorrente de saúde pública no Brasil, e Marília-SP declarou situação de "
    "emergência por epidemia da doença em fevereiro de 2025. Embora existam dados abertos de vigilância "
    "de arboviroses, eles são pouco acessíveis ao cidadão e a gestores sem formação técnica. Este trabalho "
    "tem como objetivo desenvolver um painel web acessível que disponibilize indicadores epidemiológicos da "
    "dengue em Marília e analise a relação entre variáveis climáticas e incidência. A metodologia combinou "
    "pesquisa documental, análise exploratória de dados e desenvolvimento de software orientado a testes. "
    "A solução consome a API pública do InfoDengue, armazena a série semanal em PostgreSQL, expõe uma API "
    "REST própria construída com Express e apresenta gráficos acompanhados de tabelas equivalentes para "
    "leitores de tela. Sobre 90 semanas epidemiológicas reais (2025 a 2026) foram registrados 36.937 casos, "
    "com pico de 2.925 casos na semana 10 de 2025 e 24,4% do período em nível vermelho de alerta. "
    "Observou-se correlação positiva moderada entre temperatura e incidência, mais forte para a temperatura "
    "mínima (r = 0,474), sem defasagem preferencial distinguível entre 0 e 6 semanas. O protótipo funciona "
    "ponta a ponta com dados reais e conta com 83 testes automatizados. Conclui-se, de forma preliminar, que "
    "o painel torna o dado aberto legível e que a relação clima–incidência deve ser apresentada como "
    "associação descritiva, e não como previsão."
)
trocar_texto(filhos[86], RESUMO)
trocar_texto(filhos[88], "PALAVRAS-CHAVE: Dengue; Dados abertos; Visualização de dados; Acessibilidade web; "
                         "Correlação climática.")

# Listas de ilustracoes e tabelas passam a ser campos gerados pelo Word.
trocar_texto(filhos[91], "LISTA DE ILUSTRAÇÕES")
trocar_texto(filhos[97], "LISTA DE TABELAS")
for i in (93, 94, 99, 100):
    corpo.remove(filhos[i])
for i, rotulo in ((92, "Figura"), (98, "Tabela")):
    for r in filhos[i].findall(qn("w:r")):
        filhos[i].remove(r)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f' TOC \\h \\z \\c "{rotulo}" ')
    fld.set(qn("w:dirty"), "true")
    filhos[i].append(fld)
trocar_texto(filhos[104], "")

# Quebra de secao apos o sumario: a numeracao aparece a partir da Introducao.
sect_final = corpo.find(qn("w:sectPr"))
sect_pre = copy.deepcopy(sect_final)
for ref in sect_pre.findall(qn("w:headerReference")) + sect_pre.findall(qn("w:footerReference")):
    sect_pre.remove(ref)
ppr = filhos[116].find(qn("w:pPr"))
if ppr is None:
    ppr = OxmlElement("w:pPr"); filhos[116].insert(0, ppr)
ppr.append(sect_pre)

# Remove o corpo explicativo do modelo (da Introducao ate Apendices).
for el in filhos[117:-1]:
    corpo.remove(el)

# ------------------------------------------------------------ 1 INTRODUCAO
P("1 Introdução", H1)
P("A dengue é uma arbovirose transmitida principalmente pelo mosquito *Aedes aegypti* e permanece como "
  "um dos problemas mais persistentes de saúde pública no Brasil. Sua ocorrência varia com as estações do "
  "ano: a revisão sistemática de Viana e Ignotti (2013), que reuniu 31 estudos brasileiros, concluiu que a "
  "incidência está fortemente relacionada a variáveis meteorológicas, com influência da temperatura e da "
  "pluviosidade sobre a dinâmica do vetor em todas as regiões do país.")
P("O município de Marília, no centro-oeste paulista, com cerca de 238 mil habitantes, viveu esse "
  "fenômeno de forma aguda. Em 19 de fevereiro de 2025, a Prefeitura publicou o Decreto nº 14.595, "
  "declarando situação de emergência por epidemia de dengue e instituindo um Centro de Operações de "
  "Emergência (MARÍLIA, 2025a). Quatro meses depois, o município informava ter reduzido as notificações "
  "em 73% e retornado ao alerta verde do sistema InfoDengue (MARÍLIA, 2025b). A velocidade dessa virada "
  "mostra que o acompanhamento semanal dos dados é decisivo para a percepção de risco.")
P("Os dados necessários para esse acompanhamento já são públicos. O InfoDengue, sistema desenvolvido pela "
  "Fundação Oswaldo Cruz e pela Fundação Getulio Vargas, publica semanalmente, por município, o número de "
  "casos, a incidência, o número reprodutivo efetivo (Rt), o nível de alerta e variáveis climáticas "
  "(CODEÇO et al., 2018). Contudo, esses dados chegam ao cidadão em formato técnico — uma API que devolve "
  "registros em JSON — e exigem conhecimento especializado para serem interpretados.")
P("Este trabalho situa-se na interseção entre ciência de dados e desenvolvimento web. Seu objeto é um "
  "painel web público que consome os dados abertos do InfoDengue para Marília, armazena a série histórica "
  "em banco próprio, calcula indicadores epidemiológicos e analisa a relação entre clima e incidência, "
  "com atenção à acessibilidade digital. O tema atende ao enunciado da disciplina, que pede um software "
  "com framework web, banco de dados, JavaScript, nuvem, uso de API, acessibilidade, controle de versão e "
  "testes, e trata como núcleo o item opcional de análise de dados, por ser a competência central do "
  "curso de Ciência de Dados.")
P("O relatório está organizado da seguinte forma: a seção 2.1 apresenta os objetivos; a 2.2, a "
  "justificativa e a pergunta norteadora; a 2.3, a fundamentação teórica; a 2.4, a metodologia, "
  "estruturada nas etapas ouvir, criar e implementar; e a 2.5, os resultados preliminares obtidos com "
  "o protótipo funcional sobre dados reais.")

# ------------------------------------------------------- 2 DESENVOLVIMENTO
P("2 Desenvolvimento", H1)
P("2.1 Objetivos", H2)
P("**Objetivo geral:** desenvolver e disponibilizar em nuvem um painel web acessível de vigilância da "
  "dengue em Marília-SP, que transforme dados abertos em indicadores compreensíveis e analise a relação "
  "entre variáveis climáticas e a incidência da doença.")
P("**Objetivos específicos:**")
B("Levantar a série histórica semanal de dengue em Marília a partir da API pública do InfoDengue e "
  "identificar os campos relevantes e suas limitações (exploratório);",
  "Caracterizar o período analisado por meio de indicadores epidemiológicos — total de casos, semana de "
  "pico, média móvel, tendência do Rt e distribuição dos níveis de alerta (descritivo);",
  "Analisar a correlação entre temperatura, umidade e incidência, considerando defasagens de 0 a 6 "
  "semanas (explicativo);",
  "Avaliar a acessibilidade da interface segundo as diretrizes WCAG 2.1, nível AA, garantindo "
  "alternativa textual para cada gráfico (explicativo);",
  "Verificar a correção da solução por meio de testes automatizados e publicá-la em nuvem com "
  "controle de versão público.")

P("2.2 Justificativa e delimitação do problema", H2)
P("A pergunta que norteia o trabalho é: **de que forma a apresentação de dados abertos de arboviroses em "
  "um painel acessível pode apoiar a percepção de risco e a antecipação de períodos críticos de dengue "
  "em Marília?**")
P("A **relevância social** decorre do próprio histórico recente do município: uma epidemia declarada "
  "oficialmente em 2025, com óbitos confirmados, seguida de retorno ao alerta verde em poucos meses. Em "
  "um cenário assim, informação clara e atualizada ajuda moradores a entenderem quando intensificar a "
  "eliminação de criadouros e ajuda profissionais de saúde e da educação a comunicar o risco.")
P("A **relevância acadêmica** está em tratar o dado aberto com rigor estatístico. A Lei de Acesso à "
  "Informação (BRASIL, 2011) e a Política de Dados Abertos do Poder Executivo federal (BRASIL, 2016) "
  "garantem a publicação dos dados, mas publicar não é o mesmo que tornar compreensível. Traduzir a "
  "série em indicadores, sem distorcê-los, é justamente o papel de quem trabalha com ciência de dados.")
P("A **relevância para a inclusão** vem da Lei Brasileira de Inclusão (BRASIL, 2015), cujo artigo 63 "
  "torna obrigatória a acessibilidade nos sítios da internet mantidos por órgãos de governo e por "
  "empresas com sede ou representação comercial no país. Um painel público de saúde que só comunica por "
  "gráficos exclui quem usa leitor de tela; por isso, cada gráfico do painel tem uma tabela equivalente.")
P("**Delimitação.** O escopo foi dimensionado para um grupo de um único integrante: um município "
  "(Marília, código IBGE 3529005), uma doença (dengue), as semanas epidemiológicas disponíveis de 2025 "
  "e 2026 e três visualizações. Ficam deliberadamente fora de escopo a comparação entre municípios, "
  "a previsão por aprendizado de máquina, autenticação de usuários e aplicativo móvel. A opção foi "
  "reduzir a largura e preservar a profundidade: ingestão de dados reais, análise estatística, "
  "acessibilidade e testes.")

P("2.3 Fundamentação teórica", H2)
P("**Vigilância da dengue e o sistema InfoDengue.** A vigilância epidemiológica depende de dados "
  "oportunos, mas o fluxo de notificação tem atrasos que dificultam respostas rápidas a surtos. O "
  "InfoDengue foi criado para contornar esse problema: integra dados de notificação, clima e outras "
  "fontes em um pipeline semanal, corrige estatisticamente o atraso de notificação (*nowcasting*), "
  "estima o número reprodutivo efetivo e classifica cada município em níveis de alerta (CODEÇO et al., "
  "2018). Os quatro níveis têm definição própria: verde indica condições desfavoráveis à transmissão; "
  "amarelo, condições favoráveis; laranja, transmissão sustentada; e vermelho, alta incidência. O painel "
  "utiliza exatamente essas definições para não atribuir aos dados um significado que a fonte não dá.")
P("O número reprodutivo efetivo (Rt) representa o número médio de casos secundários gerados por um caso "
  "no tempo *t*; valores acima de 1 indicam expansão da transmissão. O InfoDengue considera transmissão "
  "sustentada quando o Rt permanece acima de 1 por mais de três semanas. O painel calcula, além disso, "
  "um indicador próprio de **tendência de curto prazo**, que classifica a transmissão como em crescimento "
  "quando as duas últimas semanas têm Rt acima de 1, e em retração quando ambas estão abaixo de 1. "
  "Os dois critérios têm propósitos diferentes e não devem ser confundidos: o da fonte é um critério de "
  "alerta; o do painel, um sinal de direção.")
P("**Clima e incidência.** A literatura brasileira documenta a associação entre clima e dengue e, em "
  "especial, o seu caráter defasado: o clima de uma semana afeta o ciclo do mosquito e só depois aparece "
  "nos casos notificados. Em Foz do Iguaçu, entre 2006 e 2016, Meira et al. (2021) encontraram "
  "correlação positiva da incidência com a umidade relativa média (r = 0,276), com a temperatura média "
  "de até dois meses antes (r = 0,288) e com a pluviosidade de até três meses antes (r = 0,308). Em "
  "Uberlândia, Leite e Mendes (2023) associaram o aumento de casos à temperatura mínima acima de 18 °C, "
  "à temperatura média entre 26 °C e 28 °C e à umidade acima de 70%. Esses estudos justificam duas "
  "escolhas do projeto: testar defasagens em vez de correlacionar apenas a mesma semana, e analisar a "
  "temperatura mínima separadamente da média e da máxima. Também fornecem uma régua: coeficientes em "
  "torno de 0,3 são o que se encontra na literatura, e valores muito superiores pediriam desconfiança.")
P("**Correlação de Pearson e seus limites.** O coeficiente de Pearson mede a intensidade da associação "
  "linear entre duas variáveis, variando de −1 a 1 (MORETTIN; BUSSAB, 2017). Em séries temporais, porém, "
  "as observações não são independentes: a incidência de uma semana depende da anterior, e ambas as "
  "séries compartilham a sazonalidade. Essa autocorrelação faz com que testes de significância usuais "
  "superestimem a confiança no resultado (CHATFIELD, 2003). Por isso, neste trabalho o coeficiente é "
  "usado como medida **descritiva**, e não inferencial, e nenhuma relação causal é afirmada.")
P("**Acessibilidade digital.** As Diretrizes de Acessibilidade para Conteúdo Web, versão 2.1, organizam "
  "os requisitos em quatro princípios — perceptível, operável, compreensível e robusto — e em três níveis "
  "de conformidade (W3C, 2018). Para visualizações de dados, o ponto crítico é o critério de conteúdo "
  "não textual: um gráfico desenhado em *canvas* não é lido por tecnologias assistivas, e a informação "
  "que ele carrega precisa estar disponível também em forma textual. Além disso, a informação não pode "
  "depender apenas de cor, o que importa diretamente para um sistema de alerta baseado em verde, "
  "amarelo, laranja e vermelho.")

P("2.4 Metodologia", H2)
P("O trabalho seguiu as três etapas propostas pela Univesp — ouvir, criar e implementar — com natureza "
  "aplicada e abordagem quantitativa. Em cada etapa, as decisões foram registradas no repositório de "
  "código, o que permite rastrear a evolução real do projeto.")
P("**Ouvir e interpretar o contexto.** O contexto foi levantado por **pesquisa documental**, sem coleta "
  "com participantes até esta etapa. Foram consultadas três categorias de fonte: (a) comunicados oficiais "
  "da Prefeitura de Marília sobre a epidemia de 2025, que revelaram a situação de emergência e a rápida "
  "reversão para o alerta verde; (b) artigos científicos sobre a relação entre clima e dengue no Brasil, "
  "localizados na base SciELO e em periódicos nacionais; e (c) a documentação e a publicação "
  "metodológica do InfoDengue. Em seguida, realizou-se uma **análise exploratória** da resposta real da "
  "API para Marília, que confirmou 90 semanas epidemiológicas (da semana 1 de 2025 à semana 37 de 2026), "
  "30 campos por registro e ocorrência dos quatro níveis de alerta. Essa exploração corrigiu uma "
  "suposição inicial: esperava-se que as semanas mais recentes chegassem sem número de casos, mas nenhuma "
  "das 90 semanas veio com o dado ausente. O tratamento de semanas incompletas foi mantido como medida "
  "defensiva, e não como fenômeno observado. Para a próxima etapa, está prevista uma consulta a "
  "moradores e a profissionais de saúde, com o questionário do Apêndice A, a fim de avaliar se o painel "
  "é compreendido pelo público a que se destina.")
P("**Criar e prototipar.** Com o contexto definido, foi elaborado um documento de projeto que mapeia "
  "cada exigência do enunciado para uma decisão técnica (Tabela 1) e define a arquitetura em cinco "
  "componentes (Figura 1). A análise de dados foi isolada em funções puras — média móvel de três semanas, "
  "tendência do Rt, correlação de Pearson com defasagem e resumo dos níveis de alerta — o que permite "
  "testá-las contra valores calculados manualmente. Uma regra de dados orientou todo o projeto: **dado "
  "ausente nunca vira zero**, porque zero significa \"nenhum caso notificado\", e confundir os dois "
  "falsearia a série exatamente no ponto mais visível do painel.")
tabela("Mapeamento das exigências do enunciado para a solução",
       ["Exigência do enunciado", "Como é atendida no projeto"],
       [["Framework web", "Express (Node.js), servindo a API REST e a interface"],
        ["Banco de dados", "PostgreSQL com a série semanal e o registro de cada ingestão"],
        ["Script web (JavaScript)", "JavaScript no navegador consome a API e desenha os gráficos (Chart.js)"],
        ["Nuvem", "Aplicação no Render e banco no Neon, em planos gratuitos"],
        ["Uso de API", "Consome a API do InfoDengue e expõe API REST própria"],
        ["Acessibilidade", "WCAG 2.1 AA como meta; tabela equivalente para cada gráfico; axe-core nos testes"],
        ["Controle de versão", "Git desde o primeiro dia, com publicação no GitHub"],
        ["Testes", "Vitest (unidade), Supertest (integração) e axe-core (acessibilidade)"],
        ["Análise de dados (opcional)", "Média móvel, tendência do Rt e correlação defasada clima × incidência"]],
       [4.5, 11.0])
figura("fig-arquitetura.png", "Arquitetura da solução", 16)
P("**Implementar e testar.** O desenvolvimento seguiu a prática de escrever o teste antes do código: "
  "para cada função, primeiro se escreveu um teste que falhava e só então a implementação que o fazia "
  "passar. A ingestão é idempotente — executá-la duas vezes não duplica semanas — e roda na "
  "inicialização do servidor e depois diariamente. Cada etapa concluída passou por revisão de código, e "
  "os problemas encontrados foram corrigidos com um teste que reproduzia o defeito. Entre as correções "
  "registradas estão um deslocamento de fuso horário que alterava a data de início da semana, o "
  "armazenamento em cache de respostas de erro e a possibilidade de uma semana sem dado ser somada como "
  "zero no total do período. A estratégia de testes está resumida na Tabela 2.")
tabela("Estratégia de testes automatizados",
       ["Camada", "Ferramenta", "O que é verificado"],
       [["Análise de dados", "Vitest", "Média móvel, tendência do Rt, Pearson com defasagem e níveis, contra valores calculados à mão"],
        ["Cliente da API externa", "Vitest", "Conversão da resposta real salva em arquivo, sem depender da API estar no ar"],
        ["Persistência", "Vitest + PostgreSQL", "Gravação idempotente, conversão de tipos e datas"],
        ["API REST", "Supertest", "Formato das respostas, filtros, erro 400 e indisponibilidade do banco (503)"],
        ["Interface", "jsdom + axe-core", "Estrutura semântica, tabelas equivalentes e regras WCAG A e AA"]],
       [3.5, 3.5, 8.5])

P("2.5 Resultados preliminares: solução inicial", H2)
P("O resultado desta etapa é um **protótipo funcional ponta a ponta**, que busca os dados reais do "
  "InfoDengue, grava a série no banco, calcula os indicadores e apresenta o painel no navegador. Os "
  "resultados de cada etapa da metodologia são descritos a seguir.")
P("**Resultado da etapa ouvir.** A pesquisa documental produziu a delimitação do problema e a "
  "fundamentação para o método de análise. A análise exploratória confirmou a viabilidade técnica da "
  "fonte e corrigiu a hipótese sobre semanas incompletas, ajuste que foi levado ao código e à "
  "documentação antes de chegar a este relatório.")
P("**Resultado da etapa criar.** A Figura 2 mostra a parte superior do painel: o cabeçalho informa a data "
  "da última atualização e o número de semanas em base; os indicadores resumem o período; e o gráfico "
  "de linha apresenta os casos por semana com a média móvel de três semanas. O botão \"Ver dados em "
  "tabela\" abre a tabela equivalente ao gráfico, destinada a quem usa leitor de tela.")
figura("fig-painel-topo.png", "Painel: indicadores do período e série semanal de casos", 14.5)
P("Os indicadores obtidos sobre as 90 semanas reais estão na Tabela 3. O período foi marcado por um "
  "surto expressivo no início de 2025, coerente com a emergência declarada pelo município, seguido de "
  "longa fase de baixa transmissão.")
tabela("Indicadores do período analisado (semanas 1/2025 a 37/2026)",
       ["Indicador", "Valor"],
       [["Semanas epidemiológicas analisadas", "90"],
        ["População de referência", "238.605 habitantes"],
        ["Casos notificados no período", "36.937"],
        ["Semana de pico", "Semana 10 de 2025"],
        ["Casos na semana de pico", "2.925"],
        ["Incidência na semana de pico", "1.225,9 por 100 mil habitantes"],
        ["Rt na última semana disponível", "0,993"],
        ["Tendência de curto prazo", "Retração"]],
       [9.0, 6.5], "Fonte: elaborado pelo autor (2026), com dados do InfoDengue.")
P("A distribuição dos níveis de alerta (Figura 3 e Tabela 4) mostra que quase um quarto do período "
  "esteve em nível vermelho. No painel, as cores dos níveis sempre vêm acompanhadas do rótulo textual, "
  "de modo que a informação não depende da percepção de cor.")
figura("fig-painel-alertas.png", "Painel: distribuição dos níveis de alerta, com tabela equivalente", 14.5)
tabela("Semanas em cada nível de alerta do InfoDengue",
       ["Nível de alerta", "Semanas", "Percentual"],
       [["Verde – condições desfavoráveis à transmissão", "62", "68,9%"],
        ["Amarelo – condições favoráveis à transmissão", "1", "1,1%"],
        ["Laranja – transmissão sustentada", "5", "5,6%"],
        ["Vermelho – alta incidência", "22", "24,4%"]],
       [10.0, 2.5, 3.0], "Fonte: elaborado pelo autor (2026), com dados do InfoDengue.")
P("A análise entre clima e incidência calculou o coeficiente de Pearson entre a variável climática da "
  "semana *i* e a incidência da semana *i + k*, para *k* de 0 a 6 semanas. A Tabela 5 resume o maior "
  "coeficiente de cada variável e a Figura 4 apresenta a curva completa.")
tabela("Maior correlação entre variável climática e incidência",
       ["Variável", "Maior r", "Defasagem do maior r", "Pares (n)"],
       [["Temperatura mínima", "0,474", "3 semanas", "87"],
        ["Temperatura média", "0,426", "3 semanas", "87"],
        ["Temperatura máxima", "0,322", "0 semana", "90"],
        ["Umidade média", "0,222", "4 semanas", "86"]],
       [5.0, 2.5, 4.5, 2.5], "Fonte: elaborado pelo autor (2026), com dados do InfoDengue.")
figura("fig-correlacao-defasagem.png", "Coeficiente de correlação por defasagem entre clima e incidência", 15)
P("A leitura desses números exige cuidado. Os dados sustentam que existe **correlação positiva "
  "moderada** entre temperatura e incidência em Marília, com coeficientes entre 0,32 e 0,47 conforme a "
  "medida de temperatura. A temperatura mínima se associa mais fortemente do que a máxima, resultado "
  "convergente com Leite e Mendes (2023), e a ordem de grandeza é compatível com a encontrada por Meira "
  "et al. (2021), o que indica um resultado plausível, e não um artefato de cálculo.")
P("Os dados **não** sustentam, porém, uma defasagem preferencial. Para a temperatura média, o coeficiente "
  "vai de 0,420 na defasagem 0 a 0,426 na defasagem 3 — diferença de 0,006, desprezível em uma série de "
  "87 pares. A formulação correta é que a correlação se mantém estável entre 0 e 3 semanas e diminui "
  "levemente a partir daí. Essa conclusão levou a uma melhoria no próprio painel: o texto da seção de "
  "clima (Figura 5) passou a informar a faixa de variação do coeficiente entre as defasagens, e não "
  "apenas o valor máximo, para não sugerir ao leitor um atraso que o dado não demonstra.")
figura("fig-painel-clima.png", "Painel: relação entre temperatura média e incidência", 14.5)
P("**Resultado da etapa implementar.** A solução conta com 83 testes automatizados. A cobertura das "
  "camadas de análise de dados e de API é de 100% das linhas e funções, acima da meta de 80% definida "
  "no projeto. Os testes de acessibilidade com axe-core não apontaram violações das regras A e AA sobre "
  "a estrutura da página e sobre cada tabela gerada. É preciso registrar o alcance dessa verificação: "
  "o axe-core, executado sem renderização gráfica, não avalia o contraste de cores, e por isso a "
  "conformidade completa com a WCAG 2.1 AA ainda não pode ser afirmada; o contraste será verificado no "
  "navegador na próxima etapa. O código está versionado em Git desde o início, com 31 registros "
  "(*commits*) descritivos até a data deste relatório.")
P("**Próximos passos.** Até a entrega final estão previstos: (a) publicação do repositório no GitHub e "
  "do painel no Render, com banco no Neon; (b) aplicação do questionário do Apêndice A a moradores e "
  "profissionais de saúde, para avaliar a compreensão do painel e coletar sugestões; (c) verificação de "
  "contraste e navegação por teclado no navegador; e (d) incorporação das devolutivas ao painel e ao "
  "relatório final.")

# --------------------------------------------------------------- REFERENCIAS
P("Referências", H1)
REFS = [
    "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. **NBR 6023**: informação e documentação: referências: "
    "elaboração. Rio de Janeiro: ABNT, 2018.",
    "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. **NBR 14724**: informação e documentação: trabalhos "
    "acadêmicos: apresentação. Rio de Janeiro: ABNT, 2011.",
    "BRASIL. Lei nº 12.527, de 18 de novembro de 2011. Regula o acesso a informações previsto no inciso "
    "XXXIII do art. 5º, no inciso II do § 3º do art. 37 e no § 2º do art. 216 da Constituição Federal. "
    "**Diário Oficial da União**, Brasília, DF, 18 nov. 2011.",
    "BRASIL. Lei nº 13.146, de 6 de julho de 2015. Institui a Lei Brasileira de Inclusão da Pessoa com "
    "Deficiência (Estatuto da Pessoa com Deficiência). **Diário Oficial da União**, Brasília, DF, "
    "7 jul. 2015.",
    "BRASIL. Decreto nº 8.777, de 11 de maio de 2016. Institui a Política de Dados Abertos do Poder "
    "Executivo federal. **Diário Oficial da União**, Brasília, DF, 12 maio 2016.",
    "CHATFIELD, C. **The analysis of time series**: an introduction. 6. ed. Boca Raton: Chapman & "
    "Hall/CRC, 2003.",
    "CODEÇO, C. T. et al. Infodengue: a nowcasting system for the surveillance of arboviruses in Brazil. "
    "**Revue d'Épidémiologie et de Santé Publique**, v. 66, supl. 5, p. S386, 2018. Disponível em: "
    "https://www.sciencedirect.com/science/article/abs/pii/S0398762018311088. Acesso em: 28 set. 2026.",
    "INFODENGUE. **Sistema de alerta de arboviroses**. Rio de Janeiro: Fiocruz; FGV, 2026. Disponível "
    "em: https://info.dengue.mat.br. Acesso em: 28 set. 2026.",
    "LEITE, E. S.; MENDES, P. Análise temporal da relação entre dengue e variáveis climáticas na cidade "
    "de Uberlândia – MG. **Revista Brasileira de Geografia Física**, v. 16, n. 4, p. 2254-2270, 2023. "
    "DOI: 10.26848/rbgf.v16.4.p2254-2270.",
    "MARÍLIA (SP). Prefeitura Municipal. **Prefeitura de Marília declara Situação de Emergência por "
    "epidemia de dengue e institui Centro de Operações**. Marília, 20 fev. 2025a. Disponível em: "
    "https://www.marilia.sp.gov.br/portal/noticias/0/3/14247/. Acesso em: 29 set. 2026.",
    "MARÍLIA (SP). Prefeitura Municipal. **Marília reduz em 73% notificações da dengue e entra no Alerta "
    "Verde, aponta InfoDengue**. Marília, 18 jun. 2025b. Disponível em: "
    "https://www.marilia.sp.gov.br/portal/noticias/0/3/14963/. Acesso em: 29 set. 2026.",
    "MEIRA, M. C. R. et al. Influência do clima na ocorrência de dengue em um município brasileiro de "
    "tríplice fronteira. **Cogitare Enfermagem**, v. 26, e76974, 2021. DOI: 10.5380/ce.v26i0.76974.",
    "MORETTIN, P. A.; BUSSAB, W. O. **Estatística básica**. 9. ed. São Paulo: Saraiva, 2017.",
    "VIANA, D. V.; IGNOTTI, E. A ocorrência da dengue e variações meteorológicas no Brasil: revisão "
    "sistemática. **Revista Brasileira de Epidemiologia**, v. 16, n. 2, p. 240-256, 2013. "
    "DOI: 10.1590/S1415-790X2013000200002.",
    "W3C – WORLD WIDE WEB CONSORTIUM. **Web Content Accessibility Guidelines (WCAG) 2.1**. W3C "
    "Recommendation, 5 jun. 2018. Disponível em: https://www.w3.org/TR/WCAG21/. Acesso em: 28 set. 2026.",
]
for ref in REFS:
    P(ref, RESREF)

# ---------------------------------------------------------------- APENDICE
P("Apêndice A – Questionário de avaliação do painel", H1)
P("Instrumento a ser aplicado na próxima etapa, de forma anônima e voluntária, a moradores de Marília e "
  "a profissionais de saúde, após o participante navegar pelo painel por cerca de cinco minutos. Não "
  "serão coletados nome, contato ou qualquer dado que identifique o participante.")
B("Perfil: ( ) morador(a) ( ) profissional de saúde ( ) profissional de educação ( ) outro;",
  "Antes de hoje, você sabia que existem dados públicos semanais sobre dengue em Marília? ( ) sim ( ) não;",
  "Olhando o painel, qual é o nível de alerta atual da cidade? (resposta aberta — verifica compreensão);",
  "Em que período do ano de 2025 a dengue esteve mais alta? (resposta aberta — verifica leitura do gráfico);",
  "Em uma escala de 1 (muito difícil) a 5 (muito fácil), quão fácil foi entender os indicadores?",
  "Você usa algum recurso de acessibilidade (leitor de tela, ampliação, alto contraste)? "
  "Se sim, conseguiu acessar os dados?",
  "A frase \"correlação não estabelece causalidade\" ficou clara para você? ( ) sim ( ) não ( ) em parte;",
  "O que você mudaria ou acrescentaria no painel? (resposta aberta).")

# ------------------------------------------------ numero de pagina no cabecalho
secao = d.sections[-1]
secao.header.is_linked_to_previous = False
cab = secao.header.paragraphs[0]
cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
campo(cab, " PAGE ")
for r in cab.runs:
    r.font.name = "Times New Roman"; r.font.size = Pt(10)
d.sections[0].header.is_linked_to_previous = True

d.core_properties.author = AUTOR
d.core_properties.title = TITULO
d.save(SAIDA)
print("salvo:", SAIDA)
