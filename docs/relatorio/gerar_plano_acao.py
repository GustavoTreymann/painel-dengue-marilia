"""Preenche o modelo de Plano de Acao da Univesp (envio em atraso, Q2).
Uso (na raiz do repositorio): python docs/relatorio/gerar_plano_acao.py"""
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

RAIZ = Path(__file__).resolve().parents[2]
MODELO = RAIZ / "docs" / "modelos" / "Modelo_Plano_de_Acao.docx"
SAIDA = RAIZ / "docs" / "relatorio" / "Plano_de_Acao_PI2.docx"
AUTOR = "Luiz Gustavo Jacobuce Treymann"

d = docx.Document(MODELO)
t = d.tables


def escrever(celula, texto, negrito=False):
    p = celula.add_paragraph() if celula.text.strip() else celula.paragraphs[0]
    r = p.add_run(texto)
    r.bold = negrito
    r.font.size = Pt(11)


def preencher(sdt, texto):
    """Troca o texto-guia de um controle de conteudo pelo texto definitivo."""
    pr = sdt.find(qn("w:sdtPr"))
    for marca in pr.findall(qn("w:showingPlcHdr")):
        pr.remove(marca)
    runs = list(sdt.find(qn("w:sdtContent")).iter(qn("w:r")))
    for r in runs[1:]:
        r.getparent().remove(r)
    r = runs[0]
    rpr = r.find(qn("w:rPr"))
    if rpr is not None:
        r.remove(rpr)
    for t_el in r.findall(qn("w:t")):
        r.remove(t_el)
    t_el = OxmlElement("w:t"); t_el.text = texto; t_el.set(qn("xml:space"), "preserve")
    r.append(t_el)


controles = list(d.element.body.iter(qn("w:sdt")))
preencher(controles[0], f"{AUTOR} – RA: [preencher]")
for extra in controles[1:7]:  # grupo de um integrante: demais linhas de nome saem
    extra.getparent().remove(extra)

cabecalho = [
    "Projeto Integrador em Computação II (PJI240) – DRP07 – Turma 002 – A2026S2",
    "Desenvolver um software com framework web que utilize banco de dados, script web, nuvem, uso de "
    "API, acessibilidade, controle de versão e testes, com análise de dados. Tema do grupo: vigilância "
    "da dengue a partir de dados abertos.",
    "Painel web de vigilância da dengue em Marília: dados abertos, clima e acessibilidade",
    "Os dados abertos de vigilância da dengue são publicados semanalmente, mas em formato técnico, pouco "
    "acessível ao cidadão e a quem usa tecnologia assistiva. Pergunta norteadora: de que forma a "
    "apresentação desses dados em um painel acessível pode apoiar a percepção de risco e a antecipação "
    "de períodos críticos de dengue em Marília?",
    "Desenvolver e publicar em nuvem um painel web acessível que apresente indicadores epidemiológicos "
    "da dengue em Marília e analise a relação entre clima e incidência.",
    "Marília",
    "Jessica Caroline Pena Alves Da Silva",
]
for controle, valor in zip(controles[7:14], cabecalho):
    preencher(controle, valor)

respostas = [
    # escolha do local
    "O local escolhido foi o município de Marília-SP, sede do polo de apoio presencial do autor. Como o "
    "grupo tem um único integrante, a escolha priorizou uma comunidade em que o contato fosse viável e "
    "em que o tema norteador — software web com banco de dados, API e acessibilidade — pudesse gerar "
    "benefício concreto. A epidemia de dengue de 2025, que levou a Prefeitura a declarar situação de "
    "emergência em 19 de fevereiro de 2025 (Decreto nº 14.595), mostrou uma necessidade real de "
    "informação clara sobre o risco. A comunidade externa do projeto é formada pelos moradores de "
    "Marília e pelos profissionais de saúde e de educação que comunicam esse risco à população.",
    # conversa com a comunidade
    "Até a elaboração deste plano, o contato com a comunidade externa ocorreu de forma indireta, por "
    "meio dos comunicados oficiais da Prefeitura de Marília e dos dados públicos do sistema InfoDengue "
    "(Fiocruz/FGV), que registram a situação da cidade semana a semana. Essa leitura mostrou que o "
    "município passou de emergência declarada, em fevereiro de 2025, para o alerta verde, em junho do "
    "mesmo ano, e que essa informação está disponível apenas em formato técnico. A conversa direta com "
    "moradores e profissionais de saúde está prevista para a Quinzena 5, com a aplicação de um "
    "questionário anônimo após o uso do painel, para verificar se os indicadores são compreendidos e "
    "coletar sugestões de melhoria.",
    # problemas
    "Foram identificados três problemas relacionados ao tema norteador: (1) os dados abertos de "
    "vigilância existem, mas são publicados por uma API em formato JSON, inacessível para quem não tem "
    "formação técnica; (2) não há um meio simples de acompanhar a tendência da transmissão semana a "
    "semana, o que dificulta a percepção de risco pela população; e (3) painéis de dados costumam "
    "comunicar apenas por gráficos e cores, excluindo pessoas que usam leitor de tela ou têm "
    "dificuldade de distinguir cores. Os três se relacionam diretamente ao tema norteador, pois exigem "
    "consumo de API, banco de dados, interface web acessível e análise de dados.",
    # tema especifico
    "O tema específico é um painel web público de vigilância da dengue em Marília. O sistema consome a "
    "API do InfoDengue, armazena a série histórica em PostgreSQL, expõe uma API REST própria com "
    "Express e apresenta indicadores — casos, semana de pico, média móvel, tendência do número "
    "reprodutivo (Rt) e níveis de alerta — além da correlação entre temperatura, umidade e incidência. "
    "Cada gráfico tem uma tabela equivalente, e a interface segue as diretrizes WCAG 2.1. O projeto usa "
    "Git e GitHub para controle de versão, testes automatizados e publicação em nuvem, atendendo a "
    "todos os requisitos do tema norteador.",
]
for controle, texto in zip(controles[14:18], respostas):
    preencher(controle, texto)

CAL = "conforme calendário do AVA"
quinzenas = [
    [("Leitura do tema norteador e do enunciado do PI2", CAL, CAL, ""),
     ("Levantamento de temas possíveis relacionados a dados abertos de saúde", CAL, CAL, ""),
     ("Início do levantamento bibliográfico sobre dengue e clima", CAL, CAL, "Bases SciELO e periódicos nacionais")],
    [("Escolha da comunidade externa: município de Marília", CAL, CAL, "Grupo de um integrante"),
     ("Definição do problema e da pergunta norteadora", CAL, CAL, ""),
     ("Elaboração do Plano de Ação", CAL, "29/09/2026", "Envio em atraso, a combinar com a tutora")],
    [("Definição do título do trabalho", CAL, "28/09/2026", ""),
     ("Pesquisa documental: comunicados da Prefeitura e publicações do InfoDengue", "28/09/2026", "28/09/2026", "Contato indireto com o contexto local"),
     ("Análise exploratória da API do InfoDengue para Marília", "28/09/2026", "28/09/2026", "90 semanas reais confirmadas"),
     ("Documento de projeto e plano de implementação", "28/09/2026", "28/09/2026", "")],
    [("Implementação do protótipo: ingestão, banco, API REST e painel", "28/09/2026", "29/09/2026", "Desenvolvimento orientado a testes"),
     ("Testes automatizados e revisão de código", "28/09/2026", "29/09/2026", "86 testes"),
     ("Redação do Relatório Parcial", "29/09/2026", "30/09/2026", ""),
     ("Entrega do Relatório Parcial", "30/09/2026", "30/09/2026", "Prazo: 30/09, 23h59")],
    [("Publicação do código no GitHub e do painel em nuvem (Render e Neon)", "01/10/2026", "05/10/2026", ""),
     ("Aplicação do questionário a moradores e profissionais de saúde", "01/10/2026", "12/10/2026", "Anônimo, sem dados pessoais"),
     ("Verificação de contraste e navegação por teclado no navegador", "06/10/2026", "12/10/2026", "Complementa o axe-core"),
     ("Ajustes no painel a partir das sugestões recebidas", "08/10/2026", "15/10/2026", "")],
    [("Análise das respostas do questionário e dos resultados", "13/10/2026", "19/10/2026", ""),
     ("Finalização do protótipo e dos testes", "13/10/2026", "19/10/2026", ""),
     ("Roteiro e gravação do vídeo de apresentação", "20/10/2026", "01/11/2026", "")],
    [("Redação do Relatório Final", "02/11/2026", CAL, "Inclui devolutivas da comunidade"),
     ("Revisão final do vídeo e do relatório", CAL, CAL, ""),
     ("Entrega do Relatório Final e do vídeo", CAL, CAL, "")],
]
for tabela, atividades in zip(t[6:13], quinzenas):
    for linha, (atividade, inicio, fim, obs) in zip(tabela.rows[2:], atividades):
        for celula, valor in zip(linha.cells, (atividade, AUTOR, inicio, fim, obs)):
            escrever(celula, valor)

d.core_properties.author = AUTOR
d.save(SAIDA)
print("salvo:", SAIDA)
